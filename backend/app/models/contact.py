"""
Manual target-contact list. Stored in `discovery_seeds` with source="manual"
so the existing discovery stage already consumes them. Manual seeds never
expire and are reused every run (vs harvester seeds which are one-shot).
"""
from __future__ import annotations

from datetime import datetime
from typing import Any

from pydantic import BaseModel, ConfigDict, Field, HttpUrl


class ContactCreate(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    title: str | None = Field(default=None, max_length=200)
    company: str | None = Field(default=None, max_length=200)
    linkedin_url: HttpUrl | None = None
    # Optional group label. If set on creation/import, every parsed contact
    # lands in this group. Empty/None = ungrouped.
    group: str | None = Field(default=None, max_length=120)


class ContactBulkRequest(BaseModel):
    text: str = Field(
        min_length=1,
        max_length=200_000,
        description=(
            "One contact per line. Each line accepts: 'Name', 'Name, "
            "https://linkedin.com/in/slug', 'Name, Title, Company', or just a "
            "LinkedIn URL."
        ),
    )
    # When set, every parsed contact in this batch is assigned to the group.
    group: str | None = Field(default=None, max_length=120)


class ContactPublic(BaseModel):
    model_config = ConfigDict(populate_by_name=True)

    id: str = Field(alias="_id")
    name: str
    title: str | None = None
    company: str | None = None
    linkedin_url: str | None = None
    group: str | None = None
    status: str
    created_at: datetime


def contact_to_public(doc: dict[str, Any]) -> ContactPublic:
    return ContactPublic(
        _id=str(doc["_id"]),
        name=doc.get("extracted_name") or "",
        title=doc.get("extracted_title"),
        company=doc.get("extracted_company"),
        linkedin_url=doc.get("linkedin_url"),
        group=doc.get("group") or None,
        status=doc.get("status") or "pending",
        created_at=doc["created_at"],
    )


def parse_bulk(text: str) -> list[ContactCreate]:
    """Parse the freeform textarea input into ContactCreate rows."""
    out: list[ContactCreate] = []
    for raw in text.splitlines():
        line = raw.strip().rstrip(",")
        if not line or line.startswith("#"):
            continue
        # URL-only line.
        if line.lower().startswith(("http://", "https://")) and "," not in line:
            slug = _slug_from_url(line)
            out.append(ContactCreate(name=slug or line, linkedin_url=line))
            continue
        parts = [p.strip() for p in line.split(",")]
        name = parts[0]
        if not name:
            continue
        title: str | None = None
        company: str | None = None
        url: str | None = None
        for p in parts[1:]:
            if p.lower().startswith(("http://", "https://")):
                url = p
            elif title is None:
                title = p
            elif company is None:
                company = p
        try:
            out.append(
                ContactCreate(
                    name=name,
                    title=title,
                    company=company,
                    linkedin_url=url,  # type: ignore[arg-type]
                )
            )
        except (ValueError, TypeError):
            # malformed URL — keep the contact without it
            out.append(ContactCreate(name=name, title=title, company=company))
    return out


def _slug_from_url(url: str) -> str | None:
    if "/in/" not in url:
        return None
    try:
        return url.split("/in/", 1)[1].split("/", 1)[0].split("?", 1)[0]
    except (IndexError, ValueError):
        return None


# ---------------------------------------------------------------------------
# CSV / Excel file parser
# ---------------------------------------------------------------------------

_LINKEDIN_RE = r"https?://(?:www\.)?linkedin\.com/in/[\w-]+"

# Column name aliases (lowercase) → canonical field
_NAME_ALIASES = {"name", "full name", "fullname", "contact", "person", "first name", "firstname"}
_TITLE_ALIASES = {"title", "job title", "jobtitle", "role", "position", "designation"}
_COMPANY_ALIASES = {"company", "organization", "employer", "org", "company name"}
_URL_ALIASES = {"linkedin", "linkedin url", "linkedin_url", "linkedinurl", "url", "profile", "profile url", "linkedin profile"}


def _detect_columns(headers: list[str]) -> dict[str, int | None]:
    """Map canonical fields to column indices based on header names."""
    mapping: dict[str, int | None] = {
        "name": None,
        "title": None,
        "company": None,
        "linkedin_url": None,
    }
    for idx, raw in enumerate(headers):
        h = raw.strip().lower()
        if h in _NAME_ALIASES:
            mapping["name"] = idx
        elif h in _TITLE_ALIASES:
            mapping["title"] = idx
        elif h in _COMPANY_ALIASES:
            mapping["company"] = idx
        elif h in _URL_ALIASES:
            mapping["linkedin_url"] = idx
    return mapping


def _extract_linkedin_url(value: str) -> str | None:
    """Pull a LinkedIn profile URL from a cell value."""
    import re
    m = re.search(_LINKEDIN_RE, value)
    return m.group(0) if m else None


def _row_to_contact(row: list[str], col_map: dict[str, int | None]) -> ContactCreate | None:
    """Convert one spreadsheet row to a ContactCreate, or None if unparseable."""
    def _cell(field: str) -> str:
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return ""
        return (row[idx] or "").strip()

    name = _cell("name")
    title = _cell("title") or None
    company = _cell("company") or None
    url_raw = _cell("linkedin_url")

    url: str | None = None
    if url_raw:
        url = _extract_linkedin_url(url_raw)
        if not url and url_raw.startswith("http"):
            url = url_raw

    # If no explicit name column was found but we have a URL, derive name from slug
    if not name and url:
        name = _slug_from_url(url) or ""

    # Fall back: scan all cells for a LinkedIn URL if no URL column matched
    if not url:
        for cell in row:
            found = _extract_linkedin_url(cell)
            if found:
                url = found
                break

    if not name and not url:
        return None

    if not name:
        name = _slug_from_url(url) if url else ""

    if not name:
        return None

    try:
        return ContactCreate(
            name=name,
            title=title,
            company=company,
            linkedin_url=url,  # type: ignore[arg-type]
        )
    except (ValueError, TypeError):
        return ContactCreate(name=name, title=title, company=company)


def parse_csv_bytes(data: bytes) -> list[ContactCreate]:
    """Parse CSV file bytes into ContactCreate rows."""
    import csv
    import io

    text = data.decode("utf-8-sig")  # handle BOM
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []

    # Detect header row
    col_map = _detect_columns(rows[0])
    has_header = any(v is not None for v in col_map.values())

    if has_header:
        data_rows = rows[1:]
    else:
        # No recognized headers — assume: name, title, company, linkedin_url
        col_map = {"name": 0, "title": 1 if len(rows[0]) > 1 else None,
                    "company": 2 if len(rows[0]) > 2 else None,
                    "linkedin_url": 3 if len(rows[0]) > 3 else None}
        # But also check if column 1 is a URL
        if len(rows[0]) > 1 and _extract_linkedin_url(rows[0][1] or ""):
            col_map = {"name": 0, "title": None, "company": None, "linkedin_url": 1}
        data_rows = rows

    out: list[ContactCreate] = []
    for row in data_rows:
        if not any(cell.strip() for cell in row):
            continue
        contact = _row_to_contact(row, col_map)
        if contact:
            out.append(contact)
    return out


def parse_excel_bytes(data: bytes) -> list[ContactCreate]:
    """Parse Excel (.xlsx) file bytes into ContactCreate rows."""
    import io
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return []

    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append([str(cell).strip() if cell is not None else "" for cell in row])

    if not rows:
        return []

    col_map = _detect_columns(rows[0])
    has_header = any(v is not None for v in col_map.values())

    if has_header:
        data_rows = rows[1:]
    else:
        col_map = {"name": 0, "title": 1 if len(rows[0]) > 1 else None,
                    "company": 2 if len(rows[0]) > 2 else None,
                    "linkedin_url": 3 if len(rows[0]) > 3 else None}
        if len(rows[0]) > 1 and _extract_linkedin_url(rows[0][1] or ""):
            col_map = {"name": 0, "title": None, "company": None, "linkedin_url": 1}
        data_rows = rows

    out: list[ContactCreate] = []
    for row in data_rows:
        if not any(cell for cell in row):
            continue
        contact = _row_to_contact(row, col_map)
        if contact:
            out.append(contact)
    wb.close()
    return out
